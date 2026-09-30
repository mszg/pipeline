"""Synthetic report integration fixtures: no biological result is asserted."""
import csv
import gzip
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "workflow/scripts"))
from evaluation_metrics import dense_depth, read_vcf, hp_coverage_table
from evaluation_report import build_report


def put(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text)


def vcf(path, analysis, records=""):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt") as out:
        out.write("##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\t" + analysis + "\n" + records)


def fixture(root, genes=("ABO",), label="001", barcode="0007", sample="S001", sample_type="sample"):
    manifest = root / "samples.tsv"
    fields = ["sample", "gene", "amplicon", "fastq_input", "reference", "target_region", "sample_id", "barcode", "run_id", "sample_type"]
    records, source_rows = [], []
    for gene in genes:
        analysis = sample + "__" + gene
        results = root / "results"
        put(results / f"reference/{analysis}/reference.fasta", ">REF\nAACCGGTTAA\n")
        put(results / f"targets/{analysis}/{analysis}.bed", "REF\t0\t6\n")
        for i, (target, depth) in enumerate([("REF:1-4", [0,10,5,5,0,0,0,0,0,0]), ("REF:3-6", [0,0,10,20,10,0,0,0,0,0])], 1):
            amp = f"a{i}"
            unit = analysis + "__" + amp
            raw = f"raw/{gene}_{amp}.fastq.gz"
            records.append(dict(zip(fields, [sample, gene, amp, raw, "reference.fa", target, label, barcode, "0003", sample_type])))
            source_rows.append(f"{sample}\t{gene}\t{amp}\t{unit}\t{raw}\n")
            put(results / f"qc/raw/{unit}/NanoStats.txt", "Number of reads: 100.0\n")
            put(results / f"qc/alignment/{unit}/{unit}.alignment_qc.tsv", "metric\tvalue\nprimary_records\t80\nprimary_mapped\t50\nsecondary\t10000\nsupplementary\t10000\n")
            put(results / f"mapping/amplicons/{unit}/{unit}.variant.depth.tsv", "".join(f"REF\t{p}\t{d}\n" for p,d in enumerate(depth,1)))
        for kind, count in (("variant",80),("phasing",20)):
            put(results / f"mapping/genes/{analysis}/{analysis}.{kind}.flagstat.txt", f"{count} + 0 primary\n{count} + 0 primary mapped (100.00% : N/A)\n0 + 0 secondary\n0 + 0 supplementary\n")
        vcf(results / f"variants/{analysis}/{analysis}.norm.vcf.gz", analysis)
        vcf(results / f"phasing/{analysis}/{analysis}.phased.vcf.gz", analysis)
    with manifest.open("w", newline="") as out:
        w=csv.DictWriter(out,fieldnames=fields,delimiter="\t"); w.writeheader();w.writerows(records)
    put(root / "results/summary/input_manifest.tsv", "sample\tgene\tamplicon\tunit\tsource\n" + "".join(source_rows))
    cfg = {"samples": str(manifest), "workflow": {"run_variant_calling": True,"run_phasing":True,"run_haplotype_reconstruction":False,"run_consensus":False},
           "haplotypes":{"callable_min_depth":10},
           "reporting":{"qc":{"sample_types":{"sample":{"criteria":{
               "target_breadth_fraction":{"min":.5}, "variant_records":{"min":0}, "phased_fraction":{"min":1.0}}}}}}}
    return cfg


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.cfg=fixture(self.root)
        self.out=self.root/"synthetic.xlsx"

    def tearDown(self): self.tmp.cleanup()

    def build(self, strict=True):
        return build_report(self.cfg,self.root,self.out,REPO,strict=strict,synthetic=True)

    def test_zero_variants_no_heterozygotes_coverage_and_identifiers(self):
        data=self.build();m=data["genes"][0]["metrics"]
        self.assertEqual(m["target_bases"],6)
        self.assertEqual(m["mean_target_depth"],10)
        self.assertEqual(m["median_target_depth"],10)
        self.assertEqual(m["zero_coverage_bases"],2)
        self.assertEqual(m["target_breadth_fraction"],4/6)
        self.assertEqual(m["input_reads"],200)
        self.assertEqual(m["mapper_input_reads"],160)
        self.assertEqual(m["primary_mapped_reads"],100)
        self.assertEqual(m["mapping_percent"],62.5)
        self.assertEqual(m["variant_records"],0)
        self.assertIsNone(m["phased_fraction"])
        self.assertEqual(data["samples"][0]["overall_qc"],"PASS")
        wb=load_workbook(self.out)
        self.assertEqual(wb.sheetnames,["Sample_Summary","Gene_QC","Variants","Haplotypes","Run_Info"])
        self.assertEqual(wb["Variants"].max_row,1)
        summary=dict(zip([c.value for c in wb["Sample_Summary"][1]],[c.value for c in wb["Sample_Summary"][2]]))
        self.assertEqual((summary["run_id"],summary["sample_id"],summary["barcode"]),("0003","001","0007"))
        self.assertIn("SYNTHETIC",summary["report_mode"])
        wb.close()

    def test_multiallelic_order_and_missing_annotation_not_zero(self):
        analysis="S001__ABO"
        records="REF\t3\t.\tC\tT,G\t0\tPASS\t.\tGT:DP:AD:AF\t1/2:10:1,2,3:0.2,.\n"
        vcf(self.root/f"results/variants/{analysis}/{analysis}.norm.vcf.gz",analysis,records)
        self.build()
        wb=load_workbook(self.out);ws=wb["Variants"]
        r=dict(zip([c.value for c in ws[1]],[c.value for c in ws[2]]))
        self.assertEqual((r["ALT"],r["genotype"],r["allele_depths"],r["allele_fractions"]),("T,G","1/2","1,2,3","0.2,."))
        self.assertEqual(r["depth"],10);self.assertEqual(r["quality"],0)
        self.assertEqual(r["info_depth"],"NOT AVAILABLE");wb.close()

    def test_missing_gene_prevents_sample_pass_in_partial(self):
        self.cfg=fixture(self.root,genes=("ABO","KEL"))
        (self.root/"results/variants/S001__KEL/S001__KEL.norm.vcf.gz").unlink()
        data=self.build(False)
        self.assertEqual(len(data["samples"]),1);self.assertEqual(len(data["genes"]),2)
        self.assertEqual(data["genes"][0]["assessment"]["status"],"PASS")
        self.assertEqual(data["samples"][0]["overall_qc"],"NOT ASSESSED")
        self.assertIn("Missing variant calling",data["samples"][0]["reasons"])
        self.assertEqual(data["samples"][0]["execution_status"],"INCOMPLETE")
        with self.assertRaisesRegex(ValueError,"Strict report rejected"):
            self.build(True)

    def test_old_complete_files_do_not_imply_current_execution_success(self):
        data=self.build(False)
        self.assertEqual(data["genes"][0]["assessment"]["status"],"PASS")
        self.assertEqual(data["samples"][0]["overall_qc"],"NOT ASSESSED")
        self.assertIn("NOT VERIFIED",data["samples"][0]["execution_status"])

    def test_malformed_depth_exposed_and_dense_missing_not_zero(self):
        path=self.root/"results/mapping/amplicons/S001__ABO__a1/S001__ABO__a1.variant.depth.tsv"
        put(path,"REF\t1\t0\n")
        data=self.build(False)
        self.assertIn("Malformed",data["samples"][0]["reasons"])
        self.assertIn("EVIDENCE ERROR",data["samples"][0]["execution_status"])
        self.assertNotIn("mean_target_depth",data["genes"][0]["metrics"])
        with self.assertRaisesRegex(ValueError,"Strict report rejected"):
            self.build(True)

    def test_duplicate_source_files_withhold_aggregate_counts(self):
        path=self.root/"results/summary/input_manifest.tsv"
        put(path,path.read_text().replace("raw/ABO_a2.fastq.gz","raw/ABO_a1.fastq.gz"))
        data=self.build(False);m=data["genes"][0]["metrics"]
        self.assertIsNone(m["input_reads"]);self.assertIsNone(m["retained_variant_reads"])
        self.assertNotIn("mean_target_depth",m)
        self.assertIn("Repeated source FASTQ",data["samples"][0]["reasons"])

    def test_negative_control_does_not_inherit_sample_coverage_policy(self):
        self.cfg=fixture(self.root,sample_type="negative_control")
        self.cfg["reporting"]["qc"]["sample_types"]["negative_control"]={"criteria":{"retained_variant_reads":{"max":100}}}
        data=self.build()
        self.assertEqual(data["samples"][0]["overall_qc"],"PASS")
        self.assertNotIn("target_breadth_fraction",[x["metric"] for x in data["genes"][0]["assessment"]["criteria"]])

    def test_empty_control_policy_and_absent_cutoff_are_not_assessed(self):
        self.cfg=fixture(self.root,sample_type="positive_control")
        self.assertEqual(self.build()["samples"][0]["overall_qc"],"NOT ASSESSED")

    def test_control_partial_missing_ordinary_depth_does_not_add_coverage_requirement(self):
        self.cfg=fixture(self.root,sample_type="negative_control")
        self.cfg["reporting"]["qc"]["sample_types"]["negative_control"]={"criteria":{"input_reads":{"max":250}}}
        path=self.root/"results/mapping/amplicons/S001__ABO__a1/S001__ABO__a1.variant.depth.tsv"
        path.unlink()
        data=self.build(False)
        self.assertEqual(data["genes"][0]["assessment"]["status"],"PASS")
        self.assertEqual(data["samples"][0]["overall_qc"],"NOT ASSESSED")
        self.assertEqual(data["samples"][0]["execution_status"],"INCOMPLETE")

    def test_repeated_display_labels_without_barcode_stay_separate(self):
        path=self.root/"samples.tsv"
        with path.open() as h: rows=list(csv.DictReader(h,delimiter='\t'))
        for r in rows: r['sample_id']='Positive';r['barcode']=''
        second={**rows[0], 'sample':'S002','gene':'KEL'}
        with path.open('w',newline='') as h:
            w=csv.DictWriter(h,fieldnames=list(rows[0]),delimiter='\t');w.writeheader();w.writerows(rows+[second])
        data=self.build(False)
        self.assertEqual(len(data['samples']),2)
        self.assertEqual({r['workflow_sample'] for r in data['samples']},{'S001','S002'})

    def test_configured_failure_takes_precedence_over_missing_gene(self):
        self.cfg=fixture(self.root,genes=('ABO','KEL'))
        self.cfg['reporting']['qc']['sample_types']['sample']['criteria']['target_breadth_fraction']['min']=1.0
        (self.root/'results/variants/S001__KEL/S001__KEL.norm.vcf.gz').unlink()
        data=self.build(False)
        self.assertEqual(data['samples'][0]['overall_qc'],'FAIL')
        self.assertIn('ABO:',data['samples'][0]['reasons'])
        self.assertIn('Missing variant calling',data['samples'][0]['reasons'])
        self.cfg=fixture(self.root)
        self.cfg["reporting"]["qc"]["sample_types"]["sample"]["criteria"]={"target_breadth_fraction":{"min":None}}
        self.assertEqual(self.build()["samples"][0]["overall_qc"],"NOT ASSESSED")

    def test_formula_label_is_literal(self):
        self.cfg=fixture(self.root,label='=HYPERLINK("bad")')
        self.build();wb=load_workbook(self.out)
        ws=wb["Sample_Summary"];idx=[c.value for c in ws[1]].index("sample_id")+1
        self.assertEqual(ws.cell(2,idx).data_type,"s");wb.close()

    def test_negative_depth_and_invalid_allele_fraction_are_parse_errors(self):
        analysis="S001__ABO";p=self.root/f"results/variants/{analysis}/{analysis}.norm.vcf.gz"
        for dp,af in (("-1","0.1"),("10","1.1")):
            vcf(p,analysis,f"REF\t3\t.\tC\tT\t0\tPASS\t.\tGT:DP:AF\t0/1:{dp}:{af}\n")
            self.assertIn("Malformed variant calling",self.build(False)["samples"][0]["reasons"])

    def test_partial_consensus_completeness_uses_reference_masks_not_shifted_fasta(self):
        # Deliberately partial synthetic artifacts: a two-base insertion shifts
        # FASTA offsets, and four flank Ns must not be counted as target masks.
        self.cfg['workflow'].update(run_haplotype_reconstruction=True, run_consensus=True)
        analysis='S001__ABO'
        for hp in (1,2):
            put(self.root/f'results/consensus/{analysis}/{analysis}.haplotype{hp}.fasta', '>REF\nACCGNGTTNNNN\n')
            put(self.root/f'results/consensus/{analysis}/{analysis}.HP{hp}.mask.bed', 'REF\t2\t3\nREF\t6\t10\n')
        data=self.build(False)
        self.assertEqual(data['genes'][0]['metrics']['hp1_masked_target_bases'],1)
        self.assertEqual(data['genes'][0]['metrics']['hp1_target_unmasked_fraction'],5/6)
        wb=load_workbook(self.out);ws=wb['Haplotypes']
        rows=[dict(zip([c.value for c in ws[1]],[c.value for c in row])) for row in list(ws.rows)[1:]]
        self.assertEqual(rows[0]['consensus_length'],12)
        self.assertEqual(rows[0]['fasta_n_bases'],5)
        self.assertEqual(rows[0]['outside_target_masked_bases'],4)
        wb.close()

    def test_header_only_reference_is_an_evidence_error(self):
        put(self.root/'results/reference/S001__ABO/reference.fasta', '>REF\n')
        data=self.build(False)
        self.assertIn('EVIDENCE ERROR',data['samples'][0]['execution_status'])
        self.assertIn('contig without sequence',data['samples'][0]['reasons'])

    def test_missing_one_expected_target_cannot_pass_on_remaining_target(self):
        path=self.root/'samples.tsv'
        path.write_text(path.read_text().replace('REF:3-6',''))
        put(self.root/'results/targets/S001__ABO/S001__ABO.bed','REF\t0\t4\n')
        data=self.build()
        self.assertEqual(data['samples'][0]['overall_qc'],'NOT ASSESSED')
        self.assertNotIn('target_breadth_fraction',data['genes'][0]['metrics'])
        self.assertIn('expected amplicons lack intended target coordinates',data['samples'][0]['reasons'])

    def test_haplotype_coverage_rejects_inconsistent_counts_and_fractions(self):
        values={'target_positions':6,'callable_min_depth':10,'mean_depth':10,
                'min_depth':0,'callable_positions':4,'callable_fraction':0.666667,'below_callable_depth':2}
        path=self.root/'hp_coverage.tsv'
        def write(data):
            put(path,'metric\tHP1\tHP2\n'+''.join(f'{k}\t{v}\t{v}\n' for k,v in data.items()))
        write(values)
        hp_coverage_table(path)  # six-decimal rounding is legitimate
        for key,value in (('callable_fraction',1.1),('callable_fraction',0.5),('callable_positions',3),('target_positions',6.5)):
            with self.subTest(key=key,value=value):
                write({**values,key:value})
                with self.assertRaises(ValueError): hp_coverage_table(path)
        write({k:v for k,v in values.items() if k!='target_positions'})
        with self.assertRaisesRegex(ValueError,'Missing haplotype coverage'): hp_coverage_table(path)

    def test_assessment_provenance_includes_technical_identity(self):
        self.cfg=fixture(self.root,label='Positive')
        self.build();wb=load_workbook(self.out)
        ws=wb['Run_Info']
        headers=[c.value for c in ws[1]]
        rows=[dict(zip(headers,[c.value for c in row])) for row in list(ws.rows)[1:]]
        labels=[r['name'] for r in rows if r['category']=='Assessment']
        self.assertTrue(labels)
        self.assertTrue(all('S001__ABO' in label and 'run=0003' in label and 'barcode=0007' in label for label in labels))
        wb.close()


if __name__=="__main__": unittest.main()
