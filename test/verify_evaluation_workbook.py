#!/usr/bin/env python3
"""Independent checks of a complete real workbook against its source artifacts.

This is a reporting check, not genotype/phase truth validation. It deliberately
does not import any evaluation metric, policy or presentation implementation.
"""
import argparse
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path
import re
import statistics

from openpyxl import load_workbook


def tsv(path):
    with path.open() as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def bed_positions(path):
    result = set()
    for line in path.read_text().splitlines():
        if line and not line.startswith("#"):
            chrom, start, end, *_ = line.split("\t")
            result.update((chrom, p) for p in range(int(start), int(end)))
    return result


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def verify(path):
    payload = json.loads(path.with_suffix(".json").read_text())
    assert payload["mode"] == "strict" and not payload["synthetic"]
    root = Path(payload["workdir"]) / "results"
    manifest = tsv(Path(payload["manifest"]))
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["Sample_Summary", "Gene_QC", "Variants", "Haplotypes", "Run_Info"]
    rows = {}
    for sheet in wb:
        assert sheet.freeze_panes and sheet.auto_filter.ref
        assert not sheet.merged_cells.ranges
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        header = [cell.value for cell in sheet[1]]
        rows[sheet.title] = [dict(zip(header, [cell.value for cell in row])) for row in list(sheet.rows)[1:]]
        if rows[sheet.title]:
            assert len(sheet.tables) == 1
        for row in list(sheet.rows)[1:]:
            for column, cell in zip(header, row):
                if column == 'consensus_file' and cell.value != 'NOT AVAILABLE':
                    assert cell.hyperlink and cell.hyperlink.target == Path(cell.value).resolve().as_uri()
    fallback = str(payload["config"].get("reporting", {}).get("run_id") or "")
    identities = {(r.get("run_id") or fallback, r["sample"], r.get("sample_id") or r["sample"], r.get("barcode") or "") for r in manifest}
    expected_genes = {(r["sample"], r["gene"]) for r in manifest}
    assert len(rows["Sample_Summary"]) == len(identities)
    assert {(r["workflow_sample"], r["gene"]) for r in rows["Gene_QC"]} == expected_genes
    assert len(rows["Gene_QC"]) == len(expected_genes)
    for r in rows["Sample_Summary"]:
        assert (r["run_id"] if r["run_id"] != "NOT AVAILABLE" else "", r["workflow_sample"], r["sample_id"], r["barcode"] if r["barcode"] != "NOT AVAILABLE" else "") in identities
        assert r["execution_status"] == "DEPENDENCIES SATISFIED (cached outputs may be reused)"
        assert isinstance(r["sample_id"], str) and isinstance(r["barcode"], str)
    sources = {s["file"]: s for s in payload["sources"]}
    for name, source in sources.items():
        assert source["state"] == "READ", source
        p = Path(name)
        assert digest(p) == source["sha256"], name
        assert p.stat().st_size == source["bytes"] and p.stat().st_mtime_ns == source["mtime_ns"]
    checked = []
    for gene in rows["Gene_QC"]:
        analysis = gene["analysis"]
        configured = [r for r in manifest if r["sample"] == gene["workflow_sample"] and r["gene"] == gene["gene"]]
        target = bed_positions(root / f"targets/{analysis}/{analysis}.bed")
        intended = set()
        for r in configured:
            chrom, start, end = re.fullmatch(r"(.+):(\d+)-(\d+)", r["target_region"]).groups()
            intended.update((chrom, p) for p in range(int(start)-1, int(end)))
        assert target == intended
        depth = {p: 0 for p in target}
        raw, primary, mapped = 0, 0, 0
        for r in configured:
            unit = f"{analysis}__{r['amplicon']}"
            text = (root / f"qc/raw/{unit}/NanoStats.txt").read_text()
            raw += int(float(re.search(r"^Number of reads:\s*([\d,.]+)", text, re.M)[1].replace(",", "")))
            qc = {r["metric"]: r["value"] for r in tsv(root / f"qc/alignment/{unit}/{unit}.alignment_qc.tsv")}
            primary += int(qc["primary_records"]); mapped += int(qc["primary_mapped"])
            seen = set()
            for line in (root / f"mapping/amplicons/{unit}/{unit}.variant.depth.tsv").read_text().splitlines():
                chrom, pos, count = line.split("\t")
                key = chrom, int(pos)-1
                if key in target:
                    assert key not in seen
                    seen.add(key); depth[key] += int(count)
            assert seen == target
        assert gene["input_reads"] == raw and gene["mapper_input_reads"] == primary and gene["primary_mapped_reads"] == mapped
        assert abs(gene["mapping_percent"] - 100*mapped/primary) < 1e-9
        values = list(depth.values())
        assert gene["target_bases"] == len(target)
        assert abs(gene["mean_target_depth"] - statistics.mean(values)) < 1e-8
        assert gene["median_target_depth"] == statistics.median(values)
        assert gene["min_target_depth"] == min(values)
        assert gene["zero_coverage_bases"] == values.count(0)
        assert abs(gene["target_breadth_fraction"] - sum(d >= gene["assessment_depth"] for d in values)/len(target)) < 1e-12
        for kind, column in (("variant", "retained_variant_reads"), ("phasing", "phasing_reads")):
            text = (root / f"mapping/genes/{analysis}/{analysis}.{kind}.flagstat.txt").read_text()
            counts = re.search(r"^(\d+) \+ (\d+) primary mapped", text, re.M)
            assert gene[column] == sum(map(int, counts.groups()))
        with gzip.open(root / f"variants/{analysis}/{analysis}.norm.vcf.gz", "rt") as handle:
            calls = [line.rstrip().split("\t") for line in handle if not line.startswith("#")]
        variants = [r for r in rows["Variants"] if r["analysis"] == analysis]
        assert len(calls) == len(variants) == gene["variant_records"]
        for source, row in zip(calls, variants):
            assert (source[0], int(source[1]), source[3], source[4]) == (row["contig"], row["position"], row["REF"], row["ALT"])
            fmt = dict(zip(source[8].split(":"), source[9].split(":")))
            assert row["genotype"] == fmt["GT"]
            for tag, column in (("AD", "allele_depths"), ("AF", "allele_fractions")):
                actual, expected = row[column], fmt.get(tag)
                if expected in (None, "."):
                    assert actual == "NOT AVAILABLE"
                else:
                    assert [None if v == "." else float(v) for v in actual.split(",")] == [None if v == "." else float(v) for v in expected.split(",")]
        support = tsv(root / f"qc/haplotypes/{analysis}/{analysis}.variant_support.tsv")
        assert gene["accepted_variants"] == sum(r["STATUS"] == "ACCEPT" for r in support)
        assert gene["unresolved_variants"] == sum(r["STATUS"] == "UNRESOLVED" for r in support)
        with gzip.open(root / f"phasing/{analysis}/{analysis}.phased.vcf.gz", "rt") as handle:
            phased = [line.rstrip().split("\t") for line in handle if not line.startswith("#")]
        eligible, assigned, blocks = 0, 0, set()
        for r in phased:
            fmt = dict(zip(r[8].split(":"), r[9].split(":")))
            alleles = re.split(r"[/|]", fmt["GT"])
            if len(alleles) == 2 and "." not in alleles and alleles[0] != alleles[1]:
                eligible += 1
                if "|" in fmt["GT"] and fmt.get("PS") not in (None, "."):
                    assigned += 1; blocks.add((r[0], fmt["PS"]))
        assert (gene["eligible_heterozygous_records"], gene["phased_heterozygous_records"], gene["phase_block_count"]) == (eligible, assigned, len(blocks))
        tagging = {r['metric']: float(r['value']) for r in tsv(root/f'qc/haplotypes/{analysis}/{analysis}.haplotag_qc.tsv')}
        assert abs(gene['assigned_read_fraction'] - tagging['assigned_reads']/tagging['total_reads']) < 1e-12
        coverage = {r['metric']: r for r in tsv(root/f'qc/haplotypes/{analysis}/{analysis}.coverage_qc.tsv')}
        hap_rows = [r for r in rows["Haplotypes"] if r["analysis"] == analysis and r["row_type"] == "HAPLOTYPE"]
        assert len(hap_rows) == 2
        for hp, row in enumerate(hap_rows, 1):
            assert row['assigned_reads'] == tagging[f'hp{hp}_reads']
            assert row['min_depth'] == float(coverage['min_depth'][f'HP{hp}'])
            assert abs(row['mean_depth'] - float(coverage['mean_depth'][f'HP{hp}'])) < 1e-8
            mask = bed_positions(root / f"consensus/{analysis}/{analysis}.HP{hp}.mask.bed")
            seq = "".join(line.strip().upper() for line in Path(row["consensus_file"]).read_text().splitlines() if not line.startswith(">"))
            assert row["consensus_length"] == len(seq) and row["fasta_n_bases"] == seq.count("N")
            assert row["target_masked_reference_bases"] == len(mask & target)
            assert abs(row["target_unmasked_fraction"] - (1-len(mask & target)/len(target))) < 1e-12
        checked.append({"analysis": analysis, "input_reads": raw, "primary_mapped_reads": mapped,
                        "target_bases": len(target), "mean_target_depth": gene["mean_target_depth"],
                        "variant_records": len(calls), "accepted_variants": gene["accepted_variants"],
                        "unresolved_variants": gene["unresolved_variants"],
                        "hp_target_masks": [r["target_masked_reference_bases"] for r in hap_rows],
                        "phase_blocks": len(blocks), "phased_heterozygotes": assigned,
                        "assigned_read_fraction": gene['assigned_read_fraction'],
                        "qc": gene["gene_qc"]})
    summary = {"status": "PASS", "checked_utc": datetime.now(timezone.utc).isoformat(),
               "workbook": str(path.resolve()), "sha256": digest(path),
               "sheet_data_rows": {k: len(v) for k,v in rows.items()},
               "unchanged_metric_sources": len(sources), "analyses": checked,
               "scope": "Workbook/manifest coverage, formatting, literal cells, read counts, target depth, VCF annotations, support totals, FASTAs/masks and evidence hashes; not biological truth validation"}
    wb.close()
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = verify(args.workbook)
    text = json.dumps(result, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text)
    print(text)
