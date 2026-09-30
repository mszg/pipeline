"""Presentation-risk checks; scientific metric/assessment tests live separately."""

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "workflow/scripts"))
from evaluation_workbook import EXCEL_CELL_LIMIT, FREEZE_PANES, SHEET_NAMES, write_workbook


class EvaluationWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.path = self.root / "synthetic_presentation_test.xlsx"

    def write(self, entries=None, metadata=None):
        sheets = {name: [] for name in SHEET_NAMES}
        sheets.update(entries or {})
        write_workbook(self.path, sheets, metadata)
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        return workbook

    def test_identifiers_and_formula_like_external_values_are_literal(self):
        external = ["000123", "=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1:A2)"]
        workbook = self.write({"Sample_Summary": [
            {"sample_id": text, "barcode": "007", "run_id": "0005", "input_reads": 0}
            for text in external
        ]})
        sheet = workbook["Sample_Summary"]
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 7)], external)
        for row in range(2, 7):
            self.assertEqual(sheet.cell(row, 1).data_type, "s")
            self.assertEqual(sheet.cell(row, 2).value, "007")
            self.assertEqual(sheet.cell(row, 3).value, "0005")
            self.assertEqual((sheet.cell(row, 4).value, sheet.cell(row, 4).data_type), (0, "n"))
        values_only = load_workbook(self.path, data_only=True)
        self.addCleanup(values_only.close)
        self.assertEqual(values_only["Sample_Summary"]["A3"].value, "=1+1")

    def test_zero_missing_and_not_applicable_remain_distinct(self):
        workbook = self.write({"Gene_QC": [{
            "sample_id": "001", "variants": 0, "mean_depth": 12.25,
            "missing_depth": None, "phasing": "N/A", "status": "NOT ASSESSED",
        }]})
        sheet = workbook["Gene_QC"]
        self.assertEqual([cell.value for cell in sheet[2]], ["001", 0, 12.25, "NOT AVAILABLE", "N/A", "NOT ASSESSED"])
        self.assertEqual(sheet["B2"].data_type, "n")
        self.assertEqual(sheet["C2"].data_type, "n")

    def test_five_sheets_headers_filters_and_freezes_without_fake_rows(self):
        headers = {name: ["sample_id", "gene", "reason"] for name in SHEET_NAMES}
        workbook = self.write(metadata={"headers": headers})
        self.assertEqual(workbook.sheetnames, list(SHEET_NAMES))
        for name in SHEET_NAMES:
            sheet = workbook[name]
            self.assertEqual((sheet.max_row, sheet.max_column), (1, 3))
            self.assertEqual(sheet.freeze_panes, FREEZE_PANES[name])
            self.assertEqual(sheet.auto_filter.ref, "A1:C1")
            self.assertEqual(len(sheet.tables), 0)
            self.assertEqual(len(sheet.merged_cells.ranges), 0)

    def test_actual_rows_have_filterable_tables_and_status_text(self):
        statuses = ["PASS", "REVIEW", "FAIL", "NOT ASSESSED"]
        workbook = self.write({"Sample_Summary": [{"sample_id": str(index), "qc_status": status, "reason": "Evidence"} for index, status in enumerate(statuses)]})
        sheet = workbook["Sample_Summary"]
        self.assertEqual(list(sheet.tables), ["Evaluation_Sample_Summary"])
        self.assertEqual(sheet.tables["Evaluation_Sample_Summary"].ref, "A1:C5")
        self.assertEqual([sheet.cell(row, 2).value for row in range(2, 6)], statuses)
        self.assertEqual(len({sheet.cell(row, 2).fill.fgColor.rgb for row in range(2, 6)}), 4)
        self.assertTrue(sheet["C2"].alignment.wrap_text)

    def test_later_columns_are_retained_and_missing_values_are_explicit(self):
        workbook = self.write({"Variants": [
            {"sample_id": "001", "position": 1},
            {"sample_id": "002", "position": 2, "allele_depths": [0, 12, None]},
        ]})
        sheet = workbook["Variants"]
        self.assertEqual(sheet["C1"].value, "allele_depths")
        self.assertEqual(sheet["C2"].value, "NOT AVAILABLE")
        self.assertEqual(sheet["C3"].value, "[0, 12, null]")

    def test_reused_labels_are_not_merged_and_long_identifiers_are_text(self):
        workbook = self.write({"Sample_Summary": [
            {"sample_id": "Positive", "barcode": "001", "run_id": "12345678901234567890"},
            {"sample_id": "Positive", "barcode": "002", "run_id": "12345678901234567890"},
        ]})
        sheet = workbook["Sample_Summary"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual([sheet.cell(row, 2).value for row in (2, 3)], ["001", "002"])
        self.assertEqual(sheet["C2"].value, "12345678901234567890")
        self.assertEqual(sheet["C2"].data_type, "s")

    def test_safe_source_links_and_no_executable_schemes(self):
        source = self.root / "haplotype 001.fasta"
        source.write_text(">test\nACT\n")
        workbook = self.write({"Haplotypes": [
            {"consensus_file": source},
            {"consensus_file": source.name},
            {"consensus_file": "https://www.ncbi.nlm.nih.gov/nuccore/NG_006669.2"},
            {"consensus_file": "javascript:alert(1)"},
            {"consensus_file": "missing.fasta"},
            {"consensus_file": '=HYPERLINK("http://example.com","bad")'},
        ]})
        sheet = workbook["Haplotypes"]
        self.assertEqual(sheet["A2"].hyperlink.target, source.resolve().as_uri())
        self.assertEqual(sheet["A3"].hyperlink.target, source.resolve().as_uri())
        self.assertTrue(sheet["A4"].hyperlink.target.startswith("https://"))
        for row in (5, 6, 7):
            self.assertIsNone(sheet.cell(row, 1).hyperlink)
        self.assertEqual(sheet["A7"].data_type, "s")

    def test_cell_overflow_has_visible_warning_and_run_info_record(self):
        text = "Reason " + "x" * EXCEL_CELL_LIMIT
        workbook = self.write({
            "Gene_QC": [{"sample_id": "001", "reason": text}],
            "Run_Info": [{"section": "Run", "key": "Run identifier", "value": "synthetic"}],
        })
        cell = workbook["Gene_QC"]["B2"]
        self.assertEqual(len(cell.value), EXCEL_CELL_LIMIT)
        self.assertIn("[TRUNCATED:", cell.value)
        self.assertIn("cell limit", cell.comment.text)
        self.assertEqual(workbook["Run_Info"].max_row, 3)
        self.assertIn("Gene_QC!B2", workbook["Run_Info"]["C3"].value)

    def test_illegal_xml_characters_are_visible_not_silently_dropped(self):
        workbook = self.write({"Sample_Summary": [{"sample_id": "sample\x00A"}]})
        self.assertEqual(workbook["Sample_Summary"]["A2"].value, "sample\\u0000A")
        self.assertIn("Unicode escapes", workbook["Run_Info"]["B2"].value)

    def test_invalid_input_cannot_silently_replace_existing_report(self):
        self.write({"Gene_QC": [{"variants": 0}]})
        original = self.path.read_bytes()
        for value in (float("nan"), float("inf"), 12345678901234567):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.write({"Gene_QC": [{"depth": value}]})
            self.assertEqual(self.path.read_bytes(), original)

    def test_missing_sheet_and_duplicate_headers_fail_explicitly(self):
        with self.assertRaisesRegex(ValueError, "exactly"):
            write_workbook(self.path, {"Gene_QC": []})
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.write(metadata={"headers": {"Gene_QC": ["sample_id", "sample_id"]}})


if __name__ == "__main__":
    unittest.main()
