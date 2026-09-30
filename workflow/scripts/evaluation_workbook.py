"""Presentation-only Excel writer for the pipeline evaluation report.

Metric collection and assessment belong to the caller. In particular, this
module never converts an unavailable metric to zero or infers a QC status.

``sheets`` must contain the five names in ``SHEET_NAMES``. Each value is a list
of row dictionaries. Dictionary order determines column order; columns present
only in later rows are retained. Optional ``metadata`` accepts ``headers`` (a
sheet-to-column-list mapping, useful for empty sheets), ``link_base_dir`` (base
directory for relative file links; defaults to the workbook directory), and
``title``, ``subject`` and ``creator`` workbook properties. File/path/link/URL
columns link to existing local paths or explicit HTTP(S) URLs. Supply absolute
paths when the report will be written outside the analysis output directory.

All strings, including identifiers and potential Excel formulas, are written as
literal text. Strings exceeding Excel's cell limit are visibly abbreviated and
recorded in Run_Info; no long sequence should be passed to this writer.
"""

from __future__ import annotations

import json
import math
import os
import re
import tempfile
from collections.abc import Mapping
from datetime import date, datetime
from numbers import Integral, Real
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SHEET_NAMES = ("Sample_Summary", "Gene_QC", "Variants", "Haplotypes", "Run_Info")
MISSING_VALUE = "NOT AVAILABLE"
EXCEL_CELL_LIMIT = 32767
FREEZE_PANES = {
    "Sample_Summary": "D2", "Gene_QC": "E2", "Variants": "F2",
    "Haplotypes": "F2", "Run_Info": "B2",
}
STATUS_COLORS = {
    "PASS": "C6EFCE", "REVIEW": "FFEB9C", "FAIL": "FFC7CE",
    "NOT ASSESSED": "D9D9D9",
}
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")
_IDENTIFIERS = {
    "sample", "sample_id", "sample_name", "barcode", "barcode_id", "run",
    "run_id", "run_name", "analysis", "analysis_id", "gene", "contig",
    "chrom", "chromosome", "reference", "reference_id", "accession",
    "phase_set", "ps", "haplotype", "haplotype_id", "amplicon", "amplicon_id",
    "unit", "unit_id", "sample_key", "genotype", "gt", "ref", "alt",
}


def _key(name):
    return str(name).strip().lower().replace(" ", "_")


def _literal(value, column, location, warnings):
    """Return an Excel-supported scalar and an optional explanatory comment."""
    if value is None:
        return MISSING_VALUE, None
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, (Mapping, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True)
    elif isinstance(value, bool):
        return value, None
    elif isinstance(value, Integral):
        # Excel retains only 15 significant decimal digits. Never silently
        # round an integer; callers should supply long identifiers as strings.
        if _key(column) in _IDENTIFIERS:
            return str(value), None
        if len(str(abs(value))) > 15:
            raise ValueError(f"{location}: integer exceeds Excel's 15-digit precision")
        return int(value), None
    elif isinstance(value, Real):
        if not math.isfinite(value):
            raise ValueError(f"{location}: non-finite numeric metric; use None for missing evidence")
        if _key(column) in _IDENTIFIERS:
            return str(value), None
        return float(value), None
    elif isinstance(value, Path):
        value = str(value)
    elif not isinstance(value, str):
        raise TypeError(f"{location}: unsupported cell type {type(value).__name__}")

    notes = []
    if _ILLEGAL_XML.search(value):
        value = _ILLEGAL_XML.sub(lambda m: "\\u%04x" % ord(m.group()), value)
        notes.append("XML-incompatible characters are displayed as literal Unicode escapes")
    if len(value) > EXCEL_CELL_LIMIT:
        original_length = len(value)
        suffix = f"\n[TRUNCATED: {original_length:,} characters; inspect the linked source file.]"
        value = value[:EXCEL_CELL_LIMIT - len(suffix)] + suffix
        notes.append(f"text of {original_length:,} characters exceeded Excel's {EXCEL_CELL_LIMIT:,}-character cell limit; visible truncation marker added")
    comment = "; ".join(notes) if notes else None
    if comment:
        warnings.append(f"{location}: {comment}.")
    return value, comment


def _hyperlink(column, value, base_dir):
    """Only link explicit safe schemes or an existing local file/directory."""
    key = _key(column)
    if isinstance(value, Path):
        value = str(value)
    if not (key.endswith(("_file", "_path", "_link", "_url")) or key in {"file", "path", "link", "url", "source_file"}):
        return None
    if not isinstance(value, str) or not value or "\n" in value or "\r" in value:
        return None
    try:
        parsed = urlsplit(value)
        if parsed.scheme.lower() in {"http", "https"}:
            return value if parsed.netloc else None
        if parsed.scheme.lower() == "file":
            if parsed.netloc not in {"", "localhost"}:
                return None
            path = Path(unquote(parsed.path))
        elif parsed.scheme or value.startswith(("\\\\", "//")):
            return None
        else:
            path = Path(value).expanduser()
            if not path.is_absolute():
                path = base_dir / path
        return path.resolve().as_uri() if path.exists() else None
    except (OSError, ValueError):
        return None


def _warning_row(columns, warning):
    """Fit writer warnings into the caller's Run_Info column vocabulary."""
    keys = {_key(column): column for column in columns}
    value_key = next((keys[k] for k in ("value", "details", "description") if k in keys), None)
    name_key = next((keys[k] for k in ("key", "metric", "name", "item") if k in keys), None)
    category_key = next((keys[k] for k in ("section", "category") if k in keys), None)
    if value_key:
        row = {value_key: warning}
        if name_key:
            row[name_key] = "Excel presentation warning"
        if category_key:
            row[category_key] = "Reporting"
        return row
    return {"report_warning": warning}


def _style_sheet(sheet, columns, row_count):
    sheet.freeze_panes = FREEZE_PANES[sheet.title]
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.outlinePr.summaryRight = False
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for index, column in enumerate(columns, 1):
        key = _key(column)
        values = [sheet.cell(row, index).value for row in range(2, min(sheet.max_row, 101) + 1)]
        longest = max([len(str(column))] + [min(len(str(v)), 65) for v in values if v is not None])
        if any(part in key for part in ("reason", "explanation", "warning", "definition", "limitation", "description")):
            width = 66
        elif key.endswith(("_file", "_path", "_link", "_url")) or key in {"source", "source_file"}:
            width = 55
        else:
            width = min(44, max(14, longest + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=10, color="17365D" if cell.hyperlink else "222222", underline="single" if cell.hyperlink else None)
            if isinstance(cell.value, str):
                status = cell.value.strip().upper()
                if status in STATUS_COLORS:
                    cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[status])
                    cell.font = Font(name="Calibri", size=10, bold=True, color="222222")
                elif status in {MISSING_VALUE, "N/A", "NOT APPLICABLE"}:
                    cell.font = Font(name="Calibri", size=10, italic=True, color="666666")
                cell.number_format = "@"
            elif isinstance(cell.value, float):
                cell.number_format = "0.0000"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "0"
    end = f"{get_column_letter(len(columns))}{row_count + 1}"
    sheet.auto_filter.ref = f"A1:{end}"
    # Excel tables require data rows. A genuinely empty catalogue has headers
    # and filter controls, but no fake "zero variants" record or blank table row.
    if row_count:
        table = Table(displayName=f"Evaluation_{sheet.title}", ref=f"A1:{end}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def write_workbook(path, sheets, metadata=None):
    """Atomically write the five-sheet workbook and return its absolute path.

    Empty sheets are legitimate. ``None`` means unavailable; callers should
    provide ``N/A`` or ``NOT APPLICABLE`` for inapplicable metrics. Numeric zero
    remains numeric zero. The writer does not validate scientific conclusions.
    """
    if not isinstance(sheets, Mapping) or set(sheets) != set(SHEET_NAMES):
        raise ValueError(f"Expected exactly these sheets: {', '.join(SHEET_NAMES)}")
    metadata = dict(metadata or {})
    path = Path(path).resolve()
    base_dir = Path(metadata.get("link_base_dir", path.parent)).resolve()
    configured_headers = metadata.get("headers", {})
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = str(metadata.get("title", "ONT pipeline technical evaluation"))
    workbook.properties.subject = str(metadata.get("subject", "Technical QC and available evidence; not clinical interpretation"))
    workbook.properties.creator = str(metadata.get("creator", "ONT blood-group pipeline"))
    warnings = []

    for name in SHEET_NAMES:
        sheet = workbook.create_sheet(name)
        rows = [dict(row) for row in sheets[name]]
        columns = list(configured_headers.get(name, []))
        if len(set(columns)) != len(columns):
            raise ValueError(f"{name}: duplicate configured headers")
        for row in rows:
            for column in row:
                if column not in columns:
                    columns.append(column)
        if not columns:
            columns = ["record"]
        for column in columns:
            if not isinstance(column, str) or not column or len(column) > 255 or _ILLEGAL_XML.search(column):
                raise ValueError(f"{name}: headers must be nonempty XML-safe strings of at most 255 characters")
        if len(columns) > 16384 or len(rows) > 1048575:
            raise ValueError(f"{name}: rows or columns exceed Excel worksheet limits")
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(1, index, column)
            cell.data_type = "s"
        for row_index, row in enumerate(rows, 2):
            for column_index, column in enumerate(columns, 1):
                raw = row.get(column)
                location = f"{name}!{get_column_letter(column_index)}{row_index} ({column})"
                value, comment = _literal(raw, column, location, warnings)
                cell = sheet.cell(row_index, column_index, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if comment:
                    cell.comment = Comment(comment, "Pipeline reporting")
                link = _hyperlink(column, raw, base_dir)
                if link:
                    cell.hyperlink = link
        if name == "Run_Info" and warnings:
            for warning in warnings:
                row = _warning_row(columns, warning)
                for column in row:
                    if column not in columns:
                        columns.append(column)
                        header = sheet.cell(1, len(columns), column)
                        header.data_type = "s"
                row_index = sheet.max_row + 1
                for column_index, column in enumerate(columns, 1):
                    cell = sheet.cell(row_index, column_index, row.get(column, MISSING_VALUE))
                    cell.data_type = "s"
            # Fill any new warning-only column in preceding data rows.
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.value is None:
                        cell.value, cell.data_type = MISSING_VALUE, "s"
        _style_sheet(sheet, columns, sheet.max_row - 1)

    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
    return path
